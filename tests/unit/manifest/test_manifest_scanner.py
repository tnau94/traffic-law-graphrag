"""tests/unit/manifest/test_manifest_scanner.py — Tests for scanner metadata extraction and safety gates."""

from pathlib import Path

import docx

from src.manifest.models import DocumentItemModel, DocumentManifest
from src.manifest.scanner import (
    check_manifest_reviews,
    extract_docx_metadata,
    scan_and_update_manifest,
)


def _create_sample_docx(
    file_path: Path, so_hieu: str, title: str, authority: str = "CHÍNH PHỦ"
) -> None:
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = authority
    table.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    table.cell(1, 0).text = f"Số: {so_hieu}"
    table.cell(1, 1).text = "Hà Nội, ngày 01 tháng 01 năm 2025"

    doc.add_paragraph("")
    doc.add_paragraph(title)
    doc.add_paragraph("Căn cứ Luật Giao thông đường bộ...")
    doc.save(str(file_path))


def test_extract_docx_normal_document(tmp_path: Path):
    docx_file = tmp_path / "100_2025_ND-CP.docx"
    _create_sample_docx(
        docx_file,
        so_hieu="100/2025/NĐ-CP",
        title="NGHỊ ĐỊNH QUY ĐỊNH VỀ TỔ CHỨC VẬN TẢI ĐƯỜNG BỘ",
    )

    meta = extract_docx_metadata(docx_file)
    assert meta["number"] == "100/2025/NĐ-CP"
    assert meta["type"] == "NGHI_DINH"
    assert meta["role"] == "NORMAL"
    assert meta["need_review"] is False
    assert "100_2025_ND-CP" in meta["id"]
    assert any("Nghị định 100/2025/NĐ-CP" in a for a in meta["aliases"])


def test_extract_docx_amendment_document(tmp_path: Path):
    docx_file = tmp_path / "101_2026_ND-CP.docx"
    _create_sample_docx(
        docx_file,
        so_hieu="101/2026/NĐ-CP",
        title="NGHỊ ĐỊNH SỬA ĐỔI, BỔ SUNG NGHỊ ĐỊNH 100/2025/NĐ-CP",
    )

    meta = extract_docx_metadata(docx_file)
    assert meta["number"] == "101/2026/NĐ-CP"
    assert meta["type"] == "NGHI_DINH"
    assert meta["role"] == "AMENDMENT"
    assert meta["need_review"] is True


def test_extract_docx_omnibus_document(tmp_path: Path):
    docx_file = tmp_path / "102_2026_QH15.docx"
    _create_sample_docx(
        docx_file,
        so_hieu="102/2026/QH15",
        title="LUẬT SỬA ĐỔI, BỔ SUNG MỘT SỐ LUẬT TRONG LĨNH VỰC GIAO THÔNG",
        authority="QUỐC HỘI",
    )

    meta = extract_docx_metadata(docx_file)
    assert meta["number"] == "102/2026/QH15"
    assert meta["type"] == "LUAT"
    assert meta["role"] == "OMNIBUS"
    assert meta["need_review"] is True


def test_scan_and_update_manifest_preservation(tmp_path: Path):
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    manifest_file = tmp_path / "documents.yaml"

    # Tạo 1 file docx
    docx_file = raw_dir / "100_2025_ND-CP.docx"
    _create_sample_docx(
        docx_file,
        so_hieu="100/2025/NĐ-CP",
        title="NGHỊ ĐỊNH VẬN TẢI ĐƯỜNG BỘ",
    )

    # Quét lần 1
    manifest, added = scan_and_update_manifest(raw_dir=raw_dir, manifest_path=manifest_file)
    assert len(added) == 1
    assert added[0].number == "100/2025/NĐ-CP"

    # Giả lập dev sửa tay cấu hình trong manifest
    manifest.documents[0].scope.selected_articles = ["1", "2"]
    manifest.documents[0].need_review = False
    from src.manifest.scanner import _save_manifest
    _save_manifest(manifest, manifest_file)

    # Quét lần 2: Không được ghi đè các sửa đổi của dev
    manifest_after, added_second = scan_and_update_manifest(raw_dir=raw_dir, manifest_path=manifest_file)
    assert len(added_second) == 0
    assert len(manifest_after.documents) == 1
    assert manifest_after.documents[0].scope.selected_articles == ["1", "2"]
    assert manifest_after.documents[0].need_review is False


def test_check_manifest_reviews():
    manifest = DocumentManifest(
        documents=[
            DocumentItemModel(
                id="doc1",
                number="1/2025/ND-CP",
                raw_file="doc1.docx",
                name="Doc 1",
                type="NGHI_DINH",
                role="NORMAL",
                need_review=False,
            ),
            DocumentItemModel(
                id="doc2",
                number="2/2025/ND-CP",
                raw_file="doc2.docx",
                name="Doc 2",
                type="NGHI_DINH",
                role="AMENDMENT",
                need_review=True,
            ),
        ]
    )

    pending = check_manifest_reviews(manifest)
    assert len(pending) == 1
    assert pending[0].id == "doc2"
