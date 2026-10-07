"""src/manifest/scanner.py — Scanner phát hiện văn bản mới và cập nhật manifest."""

from __future__ import annotations

import logging
import re
import sys
from pathlib import Path
from typing import Any

import docx
import yaml

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from src.manifest.loader import load_manifest
from src.manifest.models import (
    DocumentItemModel,
    DocumentManifest,
    ScopeConfigModel,
)
from src.parser.canonical_id_resolver import canonical_document_id, normalize_so_hieu

logger = logging.getLogger(__name__)

SO_HIEU_RE = re.compile(
    r"(?:Số|Số:)\s*([0-9]+/[0-9]+/[A-Za-z0-9\-_Đđ]+)", re.IGNORECASE
)
SO_HIEU_FALLBACK_RE = re.compile(
    r"\b([0-9]+/[0-9]+/(?:QH[0-9]+|NĐ-CP|ND-CP|TT-[A-Za-z0-9\-_Đđ]+))\b", re.IGNORECASE
)


def _clean_text(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split()).strip()


def extract_docx_metadata(file_path: Path) -> dict[str, Any]:
    """Reads header and initial text of a .docx legal document to extract basic metadata."""
    doc = docx.Document(str(file_path))
    
    # 1. Thu thập văn bản từ bảng đầu (thường chứa Cơ quan ban hành & Số hiệu)
    table_texts: list[str] = []
    for table in doc.tables[:3]:
        for row in table.rows[:5]:
            for cell in row.cells:
                text = _clean_text(cell.text)
                if text:
                    table_texts.append(text)

    # 2. Thu thập các đoạn văn mở đầu
    para_texts: list[str] = [
        _clean_text(p.text)
        for p in doc.paragraphs[:25]
        if _clean_text(p.text)
    ]

    all_header_texts = table_texts + para_texts

    # 3. Trích xuất số hiệu
    number: str | None = None
    for text in all_header_texts:
        m = SO_HIEU_RE.search(text)
        if m:
            number = m.group(1).replace("–", "-").replace("—", "-")
            break
        m_fallback = SO_HIEU_FALLBACK_RE.search(text)
        if m_fallback:
            number = m_fallback.group(1).replace("–", "-").replace("—", "-")
            break

    # Nếu không tìm thấy trong nội dung, thử bóc tách từ tên file
    if not number:
        file_match = re.search(r"(\d+_\d+_[A-Z0-9\-]+)", file_path.stem)
        if file_match:
            parts = file_match.group(1).split("_")
            if len(parts) >= 3:
                number = f"{parts[0]}/{parts[1]}/{parts[2]}"

    if not number:
        number = file_path.stem

    # 4. Xác định Type (LUAT, NGHI_DINH, THONG_TU)
    doc_type = "NGHI_DINH"
    num_upper = number.upper()
    if "QH" in num_upper:
        doc_type = "LUAT"
    elif "TT" in num_upper or "THONG_TU" in num_upper:
        doc_type = "THONG_TU"
    elif "ND-CP" in num_upper or "NĐ-CP" in num_upper:
        doc_type = "NGHI_DINH"

    # 5. Xác định Tên trích yếu (Trích yếu nội dung văn bản)
    name: str = ""
    # Tìm đoạn in hoa hoặc đoạn tiêu đề chính
    for p in para_texts:
        p_clean = p.strip()
        p_lower = p_clean.lower()
        if p_clean.isupper() and len(p_clean) > 8 and not p_clean.startswith("CỘNG HÒA") and not p_clean.startswith("ĐỘC LẬP"):
            if not name or len(p_clean) > len(name):
                name = p_clean
        elif p_lower.startswith(("nghị định quy định", "luật ", "thông tư quy định")):
            name = p_clean
            break

    if not name:
        name = file_path.stem.replace("_", " ")

    # 6. Chuẩn hóa ID và Aliases
    canonical_id = canonical_document_id(number) or file_path.stem.replace(".", "_")
    aliases: list[str] = [name]
    if doc_type == "LUAT":
        aliases.extend([f"Luật số {number}", number])
    elif doc_type == "NGHI_DINH":
        aliases.extend([
            f"Nghị định {number}",
            f"Nghị định số {number}",
            number,
            normalize_so_hieu(number) or number,
        ])
    elif doc_type == "THONG_TU":
        aliases.extend([f"Thông tư {number}", number])

    # Khử trùng lặp aliases nhưng giữ thứ tự
    seen: set[str] = set()
    dedup_aliases: list[str] = []
    for a in aliases:
        if a and a not in seen:
            seen.add(a)
            dedup_aliases.append(a)

    # 7. Nhận diện Role & Safety Review Flag
    name_lower = name.lower()
    is_amendment = any(kw in name_lower for kw in ("sửa đổi", "bổ sung", "bãi bỏ", "thay thế"))
    is_omnibus = (
        ("một số luật" in name_lower)
        or ("các luật" in name_lower)
        or ("một số nghị định" in name_lower)
        or ("phân định thẩm quyền" in name_lower)
    )

    if is_omnibus:
        role = "OMNIBUS"
        need_review = True
    elif is_amendment:
        role = "AMENDMENT"
        need_review = True
    else:
        role = "NORMAL"
        need_review = False

    return {
        "id": canonical_id,
        "number": number,
        "raw_file": file_path.name,
        "name": name,
        "type": doc_type,
        "role": role,
        "need_review": need_review,
        "aliases": dedup_aliases,
    }


def scan_and_update_manifest(
    raw_dir: str | Path = "data/raw",
    manifest_path: str | Path = "config/documents.yaml",
) -> tuple[DocumentManifest, list[DocumentItemModel]]:
    """Scans raw_dir for .docx files, updates manifest_path preserving all existing entries."""
    raw_path = Path(raw_dir)
    manifest_file = Path(manifest_path)

    manifest = load_manifest(manifest_file)
    existing_files = {doc.raw_file for doc in manifest.documents if doc.raw_file}
    existing_numbers = {
        normalize_so_hieu(doc.number)
        for doc in manifest.documents
        if doc.number
    }

    if not raw_path.exists():
        logger.warning("Thư mục raw_dir %s không tồn tại.", raw_dir)
        return manifest, []

    newly_added: list[DocumentItemModel] = []
    docx_files = sorted(raw_path.glob("*.docx"))

    for file_path in docx_files:
        if file_path.name in existing_files:
            continue

        meta = extract_docx_metadata(file_path)
        norm_num = normalize_so_hieu(meta["number"])
        if norm_num in existing_numbers:
            logger.info("Bỏ qua file %s vì số hiệu %s đã tồn tại trong manifest.", file_path.name, norm_num)
            continue

        new_doc = DocumentItemModel(
            id=meta["id"],
            number=meta["number"],
            raw_file=meta["raw_file"],
            name=meta["name"],
            type=meta["type"],
            role=meta["role"],
            need_review=meta["need_review"],
            aliases=meta["aliases"],
            scope=ScopeConfigModel(mode="ALL", selected_articles=[]),
        )

        manifest.documents.append(new_doc)
        newly_added.append(new_doc)
        existing_files.add(new_doc.raw_file)
        if norm_num:
            existing_numbers.add(norm_num)

    if newly_added:
        _save_manifest(manifest, manifest_file)
        logger.info("Đã cập nhật manifest %s với %d văn bản mới.", manifest_file, len(newly_added))

    return manifest, newly_added


def _save_manifest(manifest: DocumentManifest, path: Path) -> None:
    """Saves DocumentManifest to YAML with formatted comments."""
    data = manifest.model_dump(exclude_none=True)
    header = (
        "# ==============================================================================\n"
        "#  Traffic Law GraphRAG - Document Manifest\n"
        "#  File này được tự động cập nhật bởi scanner khi có file .docx mới trong data/raw.\n"
        "#  Dev kiểm tra các văn bản có 'need_review: true' trước khi chạy pipeline nạp dữ liệu.\n"
        "# ==============================================================================\n\n"
    )
    yaml_str = yaml.dump(data, allow_unicode=True, sort_keys=False, indent=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(header + yaml_str, encoding="utf-8")


def check_manifest_reviews(manifest: DocumentManifest) -> list[DocumentItemModel]:
    """Returns all documents that still have need_review: true."""
    return [doc for doc in manifest.documents if doc.need_review]


if __name__ == "__main__":
    import sys
    manifest, added = scan_and_update_manifest()
    print("--- ĐÃ QUÉT DATA/RAW ---")
    print(f"Tổng số văn bản trong manifest: {len(manifest.documents)}")
    print(f"Số văn bản mới phát hiện: {len(added)}")
    for d in added:
        print(f"  + [{d.role}] {d.number} ({d.raw_file}) - need_review={d.need_review}")

    pending = check_manifest_reviews(manifest)
    if pending:
        print(f"\n⚠️  CẢNH BÁO: Có {len(pending)} văn bản đang chờ dev kiểm tra (need_review: true):")
        for p in pending:
            print(f"  * {p.number}: role={p.role}, file={p.raw_file}")
        sys.exit(1)
    else:
        print("\n✅ Tất cả văn bản đã được xác nhận (ready for sync).")
        sys.exit(0)
